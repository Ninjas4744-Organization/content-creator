import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from PIL import Image
from pptx import Presentation

from davinci_creator.common import ROOT, ASSETS, load
from davinci_creator.deck import html_deck, pptx_deck


class RecruitmentDeckTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.doc = load(ROOT / "projects/recruitment/content.json")

    def test_rebuilt_powerpoint_has_editable_rubik_text_and_full_slide_videos(self):
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "presentation.pptx"
            pptx_deck(self.doc, output)
            presentation = Presentation(output)
            self.assertEqual(len(presentation.slides), 13)
            for index, slide in enumerate(presentation.slides):
                if index in (3, 4, 5):
                    picture = slide.shapes[0]
                    self.assertEqual((picture.left, picture.top), (0, 0))
                    self.assertEqual((picture.width, picture.height), (presentation.slide_width, presentation.slide_height))
                    targets = [relation.target_ref for relation in slide.part.rels.values() if relation.reltype.endswith("/video")]
                    self.assertEqual(targets, [f'https://www.youtube.com/watch?v={self.doc["scenes"][index]["youtube"]}'])
                else:
                    runs = [run for shape in slide.shapes if shape.has_text_frame for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.text]
                    self.assertTrue(runs, f"Slide {index+1} must contain editable text")
                    self.assertTrue(all(run.font.name == "Rubik" for run in runs))
            self.assertFalse(any(relation.reltype.endswith(("/hyperlink", "/video")) for relation in presentation.slides[2].part.rels.values()))
            with ZipFile(output) as archive:
                ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
                for index in range(1, 14):
                    slide = ET.fromstring(archive.read(f"ppt/slides/slide{index}.xml"))
                    self.assertIsNotNone(slide.find("p:transition/p:fade", ns))

    def test_team_icons_have_transparent_corners_and_visible_foreground(self):
        teams = next(scene for scene in self.doc["scenes"] if scene["id"] == "teams")
        for item in teams["items"]:
            with self.subTest(team=item["heading"]), Image.open(ASSETS / item["image"]) as image:
                alpha = image.convert("RGBA").getchannel("A")
                self.assertEqual(alpha.getextrema(), (0, 255))
                for point in ((0, 0), (image.width-1, 0), (0, image.height-1), (image.width-1, image.height-1)):
                    self.assertEqual(alpha.getpixel(point), 0)

    def test_html_uses_live_text_and_has_no_game_link_on_slide_three(self):
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "presentation.html"
            html_deck(self.doc, output)
            content = output.read_text()
            sections = content.split('<section class="slide')[1:]
            self.assertEqual(len(sections), 13)
            third = sections[2].split("</section>", 1)[0]
            self.assertIn("FIRST", third)
            self.assertIn("<p dir=", third)
            self.assertNotIn("youtube.com", third)
            self.assertNotIn("href=", third)
            self.assertEqual(content.count('data-fullscreen="true"'), 3)
            self.assertNotIn('class="source-snapshot"', content)


if __name__ == "__main__":
    unittest.main()
