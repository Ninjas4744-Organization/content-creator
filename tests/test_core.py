import base64
import copy
import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from davinci_creator.common import ROOT,ASSETS,load,validate,voice_signature,write_json
from davinci_creator.layout import elements,textbox
from davinci_creator.deck import pptx_deck,html_deck
from davinci_creator.tts import cached_voice,extract_audio,generate
from davinci_creator.video import plan
from davinci_creator.player import build_player,vtt
from davinci_creator.serve import byte_range


class CoreTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.doc=load(ROOT/"examples/community/content.json")

    def test_brand_asset_integrity(self):
        for item in json.loads((ASSETS/"manifest.json").read_text())["assets"]:
            file=ROOT/item["path"]
            self.assertEqual(hashlib.sha256(file.read_bytes()).hexdigest(),item["sha256"])

    def test_all_example_layouts_fit(self):
        for i,scene in enumerate(self.doc["scenes"]):
            self.assertTrue(elements(scene,i,len(self.doc["scenes"])))

    def test_duplicate_ids_rejected(self):
        doc=copy.deepcopy(self.doc);doc["scenes"][1]["id"]=doc["scenes"][0]["id"]
        with self.assertRaisesRegex(ValueError,"unique"):validate(doc)

    def test_invalid_quiz_rejected(self):
        doc=copy.deepcopy(self.doc);doc["scenes"][0]["quiz"]={"question":"?","options":["a","b"],"correct":4}
        with self.assertRaisesRegex(ValueError,"quiz"):validate(doc)

    def test_excessive_text_rejected(self):
        with self.assertRaisesRegex(ValueError,"Too much text"):
            textbox("מילה "+"הרבה טקסט "*200,0,0,200,70)

    def test_media_traversal_rejected(self):
        doc=copy.deepcopy(self.doc);doc["scenes"][1]["image"]="../pyproject.toml"
        with self.assertRaisesRegex(ValueError,"escapes"):validate(doc)

    def test_native_pptx_text_and_rtl(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/"deck.pptx";pptx_deck(self.doc,path)
            prs=Presentation(path)
            self.assertEqual(len(prs.slides),12)
            visible="\n".join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
            self.assertIn("Gracious Professionalism®",visible)
            self.assertIn("מבוא לקהילה בפירסט",visible)
            with ZipFile(path) as archive:
                self.assertIn(b'rtl="1"',archive.read("ppt/slides/slide1.xml"))

    def test_html_escapes_script_content(self):
        doc=copy.deepcopy(self.doc);doc["scenes"][0]["notes"]="</script><script>alert(1)</script>"
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/"deck.html";html_deck(doc,path);content=path.read_text()
            self.assertNotIn("<script>alert(1)",content)
            self.assertIn("data:image/png;base64,",content)

    def test_gemini_response_extraction(self):
        audio=b"RIFF"+b"example"*8
        response={"steps":[{"type":"model_output","content":[{"type":"audio","data":base64.b64encode(audio).decode()}]}]}
        self.assertEqual(extract_audio(response),audio)
        with self.assertRaisesRegex(ValueError,"no audio"):extract_audio({})

    def test_cached_narration_needs_no_credentials_or_network(self):
        folder=ROOT/"examples/community/voices"
        with patch("davinci_creator.tts.request_audio",side_effect=AssertionError("Unexpected network call")):
            with patch.dict("os.environ",{},clear=True):generate(self.doc,folder)

    def test_narration_change_invalidates_cache(self):
        s=copy.deepcopy(self.doc["scenes"][0]);folder=ROOT/"examples/community/voices"
        self.assertIsNotNone(cached_voice(s,self.doc,folder))
        s["narration"]+=" שינוי בתוכן."
        self.assertIsNone(cached_voice(s,self.doc,folder))

    def test_missing_cache_file_can_be_regenerated(self):
        s=self.doc["scenes"][0]
        with tempfile.TemporaryDirectory() as temp:
            folder=Path(temp);write_json(folder/"manifest.json",{s["id"]:{"file":"missing.wav","signature":voice_signature(s,self.doc,"gemini-3.8-flash-tts")}})
            self.assertIsNone(cached_voice(s,self.doc,folder))

    def test_video_rejects_missing_narration(self):
        with tempfile.TemporaryDirectory() as temp:
            with self.assertRaisesRegex(ValueError,"Missing/stale"):plan(self.doc,Path(temp))

    def test_draft_has_no_claimed_narration(self):
        tl,voices=plan(self.doc,Path("absent"),fps=12,draft=True,seconds=2)
        self.assertTrue(tl["draft"])
        self.assertEqual(tl["duration"],24)
        self.assertTrue(all(len(a)==0 for a in voices))
        self.assertTrue(all(s["voice_duration"]==0 for s in tl["scenes"]))

    def test_interactive_player_package(self):
        tl=load(ROOT/"examples/community/approved-timeline.json")
        with tempfile.TemporaryDirectory() as temp:
            out=Path(temp)/"interactive";build_player(tl,ROOT/"examples/community/approved-video.mp4",out)
            self.assertTrue((out/"video.mp4").exists())
            content=(out/"index.html").read_text()
            self.assertIn('id="chapters"',content)
            self.assertIn('id="search"',content)
            self.assertIn('id="quiz"',content)
            self.assertNotIn("__DATA__",content)
            self.assertTrue((out/"captions.vtt").read_text().startswith("WEBVTT"))

    def test_skill_mirrors_match(self):
        canonical=ROOT/".agents/skills";mirror=ROOT/".claude/skills"
        for file in canonical.rglob("*"):
            if file.is_file():self.assertEqual(file.read_bytes(),(mirror/file.relative_to(canonical)).read_bytes())

    def test_video_byte_ranges(self):
        self.assertEqual(byte_range("bytes=0-1",100),(0,1))
        self.assertEqual(byte_range("bytes=70-",100),(70,99))
        self.assertEqual(byte_range("bytes=-20",100),(80,99))
        self.assertEqual(byte_range("bytes=90-999",100),(90,99))
        for value in ["bytes=100-", "bytes=-0", "bytes=20-10", "bytes=0-1,5-8", "invalid"]:
            with self.assertRaises(ValueError):byte_range(value,100)


if __name__=="__main__":unittest.main()
