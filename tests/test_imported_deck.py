import copy
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from davinci_creator.common import ROOT, validate
from davinci_creator.deck import build, html_deck, pptx_deck


class ImportedDeckTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        (ROOT / "outputs").mkdir(exist_ok=True)
        cls.temporary = tempfile.TemporaryDirectory(prefix="imported-test-", dir=ROOT / "outputs")
        folder = Path(cls.temporary.name)
        presentation = Presentation()
        for _ in range(4):
            presentation.slides.add_slide(presentation.slide_layouts[6])
        presentation.save(folder / "source.pptx")
        Image.new("RGB", (320, 180), "navy").save(folder / "slide.jpg")
        Image.new("RGB", (20, 20), "gold").save(folder / "animation.gif")
        prefix = folder.relative_to(ROOT).as_posix()
        slides = [{"image": f"{prefix}/slide.jpg"} for _ in range(4)]
        slides[0]["links"] = [{"x": 10, "y": 10, "w": 30, "h": 10, "href": "https://example.com/lesson", "label": "Lesson"}]
        slides[0]["overlays"] = [{"image": f"{prefix}/animation.gif", "x": 10, "y": 30, "w": 20, "h": 20}]
        for index, video in enumerate(("aaaaaaaaaaa", "bbbbbbbbbbb", "ccccccccccc"), 1):
            slides[index]["youtube"] = video
        cls.doc = {"version": 1, "title": "Imported fixture", "unlisted": True, "scenes": [
            {"id": f"scene-{i}", "title": f"Slide {i}", "layout": "statement", "body": "Editable fixture"} for i in range(4)],
            "source_deck": {"powerpoint": f"{prefix}/source.pptx", "slides": slides}}
        validate(cls.doc)

    @classmethod
    def tearDownClass(cls):
        cls.temporary.cleanup()

    def test_original_editable_powerpoint_is_preserved_byte_for_byte(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "presentation.pptx"
            pptx_deck(self.doc, path)
            self.assertEqual(path.read_bytes(), (ROOT / self.doc["source_deck"]["powerpoint"]).read_bytes())

    def test_web_deck_retains_slides_videos_links_animation_and_noindex(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "presentation.html"
            html_deck(self.doc, path)
            content = path.read_text()
            self.assertEqual(content.count('<section class="slide'), 4)
            self.assertEqual(content.count('class="source-video"'), 3)
            for video in ("aaaaaaaaaaa", "bbbbbbbbbbb", "ccccccccccc"):
                self.assertIn(f'data-youtube="{video}"', content)
            self.assertIn('class="source-link"', content)
            self.assertIn("data:image/gif;base64,", content)
            self.assertTrue('name="robots" content="noindex,nofollow"' in content)
            self.assertNotIn('<iframe src="https://', content)
            self.assertIn("frame=>frame.remove()", content)

    def test_invalid_imported_sources_fail_before_building(self):
        for mutation in ("outside", "count", "video", "link", "geometry", "crop"):
            doc = copy.deepcopy(self.doc)
            slides = doc["source_deck"]["slides"]
            if mutation == "outside":
                slides[0]["image"] = "../outside.jpg"
            elif mutation == "count":
                slides.pop()
            elif mutation == "video":
                slides[1]["youtube"] = 'invalid" onload="example'
            elif mutation == "link":
                slides[0]["links"][0]["href"] = "javascript:example()"
            elif mutation == "geometry":
                slides[0]["links"][0]["x"] = float("nan")
            else:
                slides[0]["overlays"][0]["crop"] = {"top": 0.8, "bottom": 0.8}
            with self.subTest(mutation=mutation), self.assertRaises(ValueError):
                validate(doc)

    def test_source_storyboard_has_all_imported_slides(self):
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)
            build(self.doc, output)
            self.assertEqual(len(list((output / "previews").glob("slide-*.png"))), 4)
            self.assertTrue((output / "previews/storyboard.jpg").is_file())


if __name__ == "__main__":
    unittest.main()
