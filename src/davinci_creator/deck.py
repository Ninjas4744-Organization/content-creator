from __future__ import annotations

import base64
import html
import json
import mimetypes
import re
import shutil
from pathlib import Path
from .common import ASSETS, ROOT, local_file
from .layout import W,H,elements
from .render import storyboard


def data_url(name):
    path=local_file(ASSETS,name)
    mime=mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    return f"data:{mime};base64,"+base64.b64encode(path.read_bytes()).decode()


def html_deck(doc,out):
    css="""
@font-face{font-family:Rubik;src:url(REGULAR)}
@font-face{font-family:Rubik;src:url(BOLD);font-weight:700}
*{box-sizing:border-box}body{margin:0;background:#08152b;color:#fafcff;font-family:Rubik,Arial,sans-serif;overflow:hidden}
#stage{position:absolute;width:1280px;height:720px;left:50%;top:50%;transform-origin:center}
.slide{position:absolute;inset:0;display:none;background-size:cover}.slide.active{display:block}
.el{position:absolute;white-space:pre-wrap;line-height:1.25;margin:0}.slide.active .el{animation:enter .6s both;animation-delay:var(--delay)}
@keyframes enter{from{opacity:0;translate:0 18px}to{opacity:1;translate:0 0}}
nav{position:fixed;bottom:8px;left:50%;translate:-50% 0;display:flex;align-items:center;gap:12px;background:#08152be8;border-radius:20px;padding:6px 14px;opacity:.25;transition:opacity .2s}nav:hover,nav:focus-within{opacity:1}button{font:inherit;color:inherit;background:#244666;border:0;border-radius:8px;padding:7px 15px;cursor:pointer}
#count{direction:ltr;unicode-bidi:isolate}#notes{display:none;position:fixed;bottom:55px;right:5%;width:90%;padding:20px;background:#08152bf5;direction:rtl;border:1px solid #a9d8f6;border-radius:12px;max-height:35vh;overflow:auto}
.source-snapshot{position:absolute;inset:0;width:100%;height:100%;object-fit:contain}.source-overlay,.source-link{position:absolute}.source-overlay{object-fit:fill}.source-link:focus-visible{outline:4px solid #ffda9d}.source-video{position:absolute;inset:180px 120px 95px;display:grid;place-items:center}.youtube-play{display:grid;place-items:center;gap:16px;font-size:28px;background:#08152be8;border:2px solid #ffda9d;padding:24px 48px;border-radius:20px}.play-icon{font-size:64px;color:#ffda9d}.source-video iframe{position:absolute;inset:0;width:100%;height:100%;border:0;background:#000}.youtube-external{position:absolute;bottom:48px;right:80px;color:#fafcff;font-size:21px}.youtube-play[hidden]{display:none}
.slide.video-slide{background:#000}.video-slide .source-video{inset:0}.video-slide .youtube-play{position:absolute;inset:0;width:100%;height:100%;border:0;background:transparent;padding:0}.video-slide .youtube-play:focus-visible{outline:3px solid #ffda9d;outline-offset:-6px}.video-slide.playing .youtube-external{display:none}
@media(prefers-reduced-motion:reduce){.slide.active .el{animation:none}}
@media print{@page{size:13.333in 7.5in;margin:0}body{overflow:visible}#stage{position:static;transform:none!important}nav,#notes{display:none!important}.slide{position:relative;display:block!important;width:1280px;height:720px;break-after:page;print-color-adjust:exact}.el{animation:none!important}}
""".replace("REGULAR",data_url("brand/Rubik-Regular.ttf")).replace("BOLD",data_url("brand/Rubik-Bold.ttf"))
    css+=f'.slide{{background-image:url({data_url("brand/space-content.png")})}}.slide.hero{{background-image:url({data_url("brand/space-hero.png")})}}'
    sections=[]
    for i,scene in enumerate(doc["scenes"]):
        parts=[]
        if doc.get("source_deck"):
            from .imported_deck import slide_markup
            parts.append(slide_markup(doc["source_deck"]["slides"][i], scene))
        for el in ([] if doc.get("source_deck") else elements(scene,i,len(doc["scenes"]))):
            style=f'left:{el["x"]}px;top:{el["y"]}px;--delay:{el["delay"]}s;'
            if el["kind"]=="line":
                parts.append(f'<div class="el" style="{style}width:{max(1,el["x2"]-el["x"])}px;height:{max(1,el["y2"]-el["y"])}px;background:#{el["color"]};opacity:.65"></div>')
            elif el["kind"]=="image":
                crop = el.get("crop", {})
                if crop:
                    width = 100/(1-crop.get("left", 0)-crop.get("right", 0))
                    height = 100/(1-crop.get("top", 0)-crop.get("bottom", 0))
                    inner = f'position:absolute;width:{width}%;height:{height}%;left:{-crop.get("left", 0)*width}%;top:{-crop.get("top", 0)*height}%'
                    parts.append(f'<div class="el" style="{style}width:{el["w"]}px;height:{el["h"]}px;overflow:hidden"><img alt="" style="{inner}" src="{data_url(el["src"])}"></div>')
                else:
                    parts.append(f'<img class="el" alt="" style="{style}width:{el["w"]}px;height:{el["h"]}px;object-fit:contain" src="{data_url(el["src"])}">')
            elif el["kind"] == "play":
                parts.append(f'<svg class="el" aria-hidden="true" style="{style}width:{el["w"]}px;height:{el["h"]}px" viewBox="0 0 176 176"><circle cx="88" cy="88" r="85" fill="none" stroke="#{el["color"]}" stroke-width="3"/><path d="M70 49 L70 127 L128 88 Z" fill="#{el["color"]}"/></svg>')
            else:
                direction="rtl" if el["rtl"] else "ltr"
                parts.append(f'<p dir="{direction}" class="el" style="{style}width:{el["w"]}px;height:{el["h"]}px;font-size:{el["size"]}px;font-weight:{700 if el["bold"] else 400};text-align:{el["align"]};color:#{el["color"]}">{html.escape(el["text"])}</p>')
        if scene.get("layout") == "video":
            title = html.escape(scene["title"], quote=True)
            video = scene["youtube"]
            parts.append(f'<div class="source-video" data-youtube="{video}" data-title="{title}" data-fullscreen="true"><button class="youtube-play" aria-label="ניגון {title} במסך מלא"></button></div><a class="youtube-external" dir="rtl" href="https://www.youtube.com/watch?v={video}" target="_blank" rel="noopener noreferrer">פתיחה ב־<bdi dir="ltr">YouTube</bdi> ↗</a>')
        sections.append(f'<section class="slide{" active" if i==0 else ""}{" hero" if scene.get("layout")=="cover" else ""}{" video-slide" if scene.get("layout")=="video" else ""}" aria-label="{html.escape(scene["title"],quote=True)}">'+"".join(parts)+"</section>")
    notes=json.dumps([s.get("notes",s.get("narration","")) for s in doc["scenes"]],ensure_ascii=False).replace("<","\\u003c")
    js="""
const slides=[...document.querySelectorAll('.slide')],notes=NOTES;let index=0;
function resize(){document.querySelector('#stage').style.transform=`translate(-50%,-50%) scale(${Math.min(innerWidth/1280,innerHeight/720)})`}
function show(n){index=Math.max(0,Math.min(slides.length-1,Number.isFinite(n)?Math.trunc(n):0));slides.forEach((s,i)=>{s.classList.toggle('active',i===index);s.setAttribute('aria-hidden',String(i!==index));if(i!==index){s.classList.remove('playing');s.querySelectorAll('.source-video iframe').forEach(frame=>frame.remove());s.querySelectorAll('.youtube-play').forEach(button=>button.hidden=false)}});document.querySelector('#count').textContent=`${index+1} / ${slides.length}`;document.querySelector('#notes').textContent=notes[index];location.hash=String(index+1)}
document.querySelectorAll('.youtube-play').forEach(button=>{button.onclick=()=>{const box=button.closest('.source-video'),frame=document.createElement('iframe');if(box.dataset.fullscreen==='true'&&!document.fullscreenElement)document.documentElement.requestFullscreen?.().catch(()=>{});frame.src=`https://www.youtube-nocookie.com/embed/${box.dataset.youtube}?autoplay=1&rel=0`;frame.title=box.dataset.title;frame.referrerPolicy='strict-origin-when-cross-origin';frame.allow='autoplay; encrypted-media; picture-in-picture; fullscreen';frame.allowFullscreen=true;button.hidden=true;box.closest('.slide').classList.add('playing');box.append(frame)}});
document.querySelector('#next').onclick=()=>show(index+1);document.querySelector('#prev').onclick=()=>show(index-1);
function toggleNotes(){const el=document.querySelector('#notes');el.style.display=el.style.display==='block'?'none':'block'}
document.querySelector('#noteButton').onclick=toggleNotes;
document.querySelector('#full').onclick=()=>document.fullscreenElement?document.exitFullscreen():document.documentElement.requestFullscreen?.();
addEventListener('keydown',e=>{if(['ArrowLeft','PageDown',' '].includes(e.key)){e.preventDefault();show(index+1)}if(['ArrowRight','PageUp'].includes(e.key))show(index-1);if(e.key.toLowerCase()==='n')toggleNotes();if(e.key==='Home')show(0);if(e.key==='End')show(slides.length-1)});
addEventListener('hashchange',()=>show(Number(location.hash.slice(1)||1)-1));addEventListener('resize',resize);resize();show(Number(location.hash.slice(1)||1)-1);
""".replace("NOTES",notes)
    robots = '<meta name="robots" content="noindex,nofollow">' if doc.get("unlisted") else ""
    out.write_text(f'<!doctype html><html lang="he"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">{robots}<title>{html.escape(doc["title"])}</title><style>{css}</style><main id="stage">'+"".join(sections)+'</main><aside id="notes"></aside><nav dir="rtl" aria-label="ניווט במצגת"><button id="prev">הקודם</button><span id="count" aria-live="polite"></span><button id="next">הבא</button><button id="noteButton">הערות</button><button id="full">מסך מלא</button></nav><script>'+js+'</script></html>',encoding="utf-8")


def pptx_deck(doc,out):
    if doc.get("source_deck"):
        shutil.copy2(local_file(ROOT, doc["source_deck"]["powerpoint"]), out)
        return
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement
    from PIL import Image
    prs=Presentation()
    prs.slide_width=Inches(W/96)
    prs.slide_height=Inches(H/96)
    for i,scene in enumerate(doc["scenes"]):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        if scene.get("layout") == "video":
            from io import BytesIO
            from .render import frame
            from pptx.oxml.ns import qn
            poster = BytesIO()
            frame(scene, i, len(doc["scenes"])).save(poster, format="PNG")
            poster.seek(0)
            picture = slide.shapes.add_picture(poster, 0, 0, prs.slide_width, prs.slide_height)
            url = f'https://www.youtube.com/watch?v={scene["youtube"]}'
            relation = slide.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/video", is_external=True)
            video_file = OxmlElement("a:videoFile")
            video_file.set(qn("r:link"), relation)
            picture._element.nvPicPr.nvPr.append(video_file)
            picture.click_action.hyperlink.address = url
        else:
            bg=ASSETS/"brand"/("space-hero.png" if scene.get("layout")=="cover" else "space-content.png")
            slide.shapes.add_picture(str(bg),0,0,prs.slide_width,prs.slide_height)
        for el in ([] if scene.get("layout") == "video" else elements(scene,i,len(doc["scenes"]))):
            px=lambda v:Inches(v/96)
            if el["kind"]=="image":
                path=local_file(ASSETS,el["src"])
                with Image.open(path) as img:iw,ih=img.size
                crop = el.get("crop", {})
                visible_w = iw*(1-crop.get("left", 0)-crop.get("right", 0))
                visible_h = ih*(1-crop.get("top", 0)-crop.get("bottom", 0))
                ratio=min(el["w"]/visible_w,el["h"]/visible_h)
                picture = slide.shapes.add_picture(str(path),px(el["x"]+(el["w"]-visible_w*ratio)/2),px(el["y"]+(el["h"]-visible_h*ratio)/2),px(visible_w*ratio),px(visible_h*ratio))
                for side in ("left", "top", "right", "bottom"):
                    setattr(picture, f"crop_{side}", crop.get(side, 0))
            elif el["kind"]=="line":
                from pptx.enum.shapes import MSO_CONNECTOR
                shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,px(el["x"]),px(el["y"]),px(el["x2"]),px(el["y2"]))
                shape.line.color.rgb=RGBColor.from_string(el["color"])
                shape.line.width=Pt(.75)
            else:
                shape=slide.shapes.add_textbox(px(el["x"]),px(el["y"]),px(el["w"]),px(el["h"]))
                tf=shape.text_frame
                tf.clear();tf.word_wrap=False;tf.vertical_anchor=MSO_ANCHOR.TOP
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                for j,line in enumerate(el["text"].split("\n")):
                    p=tf.paragraphs[0] if j==0 else tf.add_paragraph()
                    p.alignment={"right":PP_ALIGN.RIGHT,"center":PP_ALIGN.CENTER,"left":PP_ALIGN.LEFT}[el["align"]]
                    p._p.get_or_add_pPr().set("rtl","1" if re.search(r"[\u0590-\u05ff]",line) else "0")
                    p.line_spacing=Pt(el["size"]*.75*1.25)
                    p.space_after=Pt(0);p.space_before=Pt(0)
                    run=p.add_run();run.text=line
                    run.font.name="Rubik";run.font.size=Pt(el["size"]*.75);run.font.bold=el["bold"]
                    run.font.color.rgb=RGBColor.from_string(el["color"])
                    rpr=run._r.get_or_add_rPr();rpr.set("lang","he-IL" if el["rtl"] else "en-US")
                    for tag in ["a:cs","a:ea"]:
                        node=OxmlElement(tag);node.set("typeface","Rubik");rpr.append(node)
        if doc.get("slide_transition") == "fade":
            transition = OxmlElement("p:transition")
            transition.set("spd", "med")
            transition.append(OxmlElement("p:fade"))
            slide._element.insert(2, transition)
        slide.notes_slide.notes_text_frame.text=scene.get("notes",scene.get("narration",""))
    prs.core_properties.title=doc["title"]
    prs.core_properties.author="Da Vinci 4744"
    prs.save(out)


def build(doc,out:Path):
    out.mkdir(parents=True,exist_ok=True)
    pptx_deck(doc,out/"presentation.pptx")
    html_deck(doc,out/"presentation.html")
    if doc.get("source_deck"):
        from .imported_deck import storyboard as imported_storyboard
        imported_storyboard(doc, out/"previews")
    else:
        storyboard(doc,out/"previews")
    print(f"Created {len(doc['scenes'])} slides, PowerPoint, HTML and previews: {out}")
